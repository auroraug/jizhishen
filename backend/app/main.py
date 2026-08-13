import hashlib
import io
import json
import re
import shutil
import tempfile
import zipfile
import fitz
from datetime import datetime
from difflib import SequenceMatcher
from pathlib import Path

from fastapi import BackgroundTasks, FastAPI, File, Form, HTTPException, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse, JSONResponse, Response
from pydantic import BaseModel
from pypdf import PdfReader
from docx import Document
from reportlab.lib.pagesizes import A4
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.pdfgen import canvas
from starlette.background import BackgroundTask

from .db import FILES_DIR, PARSER_ARTIFACTS_DIR, TRACES_DIR, db, decode, init_db, rows
from .audit_pipeline import CANONICAL_PHASES, RULES, canonical_phase, execute_audit, model_json
from .document_parser import locate_quote, parse_document
from .providers import discover_models, mineru_batch_result, mineru_upload_files, test_provider, test_provider_config
from .seed import DOC_TYPES, lines_for, make_pdf, seed
from .model_config import ALL_STAGES, load_config, merge_public_update, public_config, resolve_route, snapshot
from .prompt_store import create_draft, list_prompts, prompt_snapshot, publish, published_prompt, render, seed_prompts, update_draft
from .trace_store import attach_input, finish_span, load_artifact, run_trace, skipped_span, start_span
from .parser_service import (NORMALIZER_VERSION, PYMUPDF_VERSION, activate_version, create_renormalized_version, create_version, ensure_legacy_version,
    poll_mineru_batch, poll_mineru_local_task, run_pymupdf_version, run_renormalized_version, version_payload)
from .review_actions import ACTION_TYPES, REVIEW_MODES, action_definition, normalized_action, parse_cutoff

app=FastAPI(title="集智审 API",version="1.0.0-mvp",docs_url="/api/docs",openapi_url="/api/openapi.json")
app.add_middleware(CORSMiddleware,allow_origins=["*"],allow_methods=["*"],allow_headers=["*"])

@app.on_event("startup")
def startup():
    init_db(); seed_prompts(); load_config()

@app.get("/api/health")
def health():
    with db() as conn:
        counts={t:conn.execute(f"SELECT COUNT(*) FROM {t}").fetchone()[0] for t in ("projects","anomalies")}
        counts["documents"]=conn.execute("SELECT COUNT(*) FROM documents WHERE deleted_at IS NULL").fetchone()[0]
    return {"ok":True,"service":"集智审","version":"1.0.0-mvp","counts":counts,"time":datetime.now().isoformat()}

@app.get("/api/dashboard")
def dashboard():
    with db() as conn:
        current_run_sql="""a.run_id=(SELECT r.id FROM audit_runs r WHERE r.project_id=a.project_id AND r.status='completed' AND r.run_kind='audit' ORDER BY r.id DESC LIMIT 1)"""
        counts={"projects":conn.execute("SELECT COUNT(*) FROM projects WHERE is_demo=0").fetchone()[0],
                "demo_projects":conn.execute("SELECT COUNT(*) FROM projects WHERE is_demo=1").fetchone()[0],
                "documents":conn.execute("SELECT COUNT(*) FROM documents d JOIN projects p ON p.id=d.project_id WHERE p.is_demo=0 AND d.deleted_at IS NULL").fetchone()[0],
                "anomalies":conn.execute(f"SELECT COUNT(*) FROM anomalies a JOIN projects p ON p.id=a.project_id WHERE p.is_demo=0 AND a.status='待复核' AND {current_run_sql}").fetchone()[0],
                "high":conn.execute(f"SELECT COUNT(*) FROM anomalies a JOIN projects p ON p.id=a.project_id WHERE p.is_demo=0 AND a.status='待复核' AND a.severity='high' AND {current_run_sql}").fetchone()[0],
                "runs":conn.execute("SELECT COUNT(*) FROM audit_runs r JOIN projects p ON p.id=r.project_id WHERE p.is_demo=0 AND r.run_kind='audit'").fetchone()[0],
                "facts":conn.execute("SELECT COUNT(*) FROM extracted_facts f JOIN projects p ON p.id=f.project_id WHERE p.is_demo=0").fetchone()[0]}
        phases={phase:conn.execute("SELECT COUNT(*) FROM documents d JOIN projects p ON p.id=d.project_id WHERE d.doc_type=? AND p.is_demo=0 AND d.deleted_at IS NULL",(phase,)).fetchone()[0]
                for phase in CANONICAL_PHASES}
        latest=conn.execute("SELECT r.id,r.project_id,r.status,r.started_at,r.finished_at,r.fact_count,r.anomaly_count,r.provider FROM audit_runs r JOIN projects p ON p.id=r.project_id WHERE p.is_demo=0 AND r.run_kind='audit' ORDER BY r.id DESC LIMIT 1").fetchone()
    return {"counts":counts,"document_phases":phases,"latest_run":dict(latest) if latest else None,
            "connectors":[
                {"id":"manual_upload","name":"文件上传适配器","status":"active","records":counts["documents"],"mode":"真实"},
                 *[{"id":key,"name":name,"status":"pending","records":None,"mode":"尚未配置"} for key,name in
                  (("oa","OA审批系统"),("project","工程项目系统"),("assets","三资监管系统"),
                   ("finance","财务系统"),("tender","招投标系统"),("archive","档案文件服务器"))]
            ]}

@app.get("/api/rules")
def rules():
    with db() as conn:
        items=rows(conn.execute("SELECT * FROM rule_definitions ORDER BY system_managed DESC,code").fetchall())
    return {"items":items,"total":len(items)}

class RuleWrite(BaseModel):
    code: str | None = None
    name: str
    kind: str = "LLM语义规则"
    fields: str = ""
    description: str = ""
    severity: str = "medium"
    enabled: bool = True

RULE_SEVERITIES={"high","moderate","minor","medium","low"}
RULE_KINDS={"LLM语义规则","人工核验规则","确定性规则"}

def validate_rule(body: RuleWrite, creating=False):
    code=(body.code or "").strip().upper()
    if creating and not re.fullmatch(r"[A-Z][A-Z0-9_-]{2,31}",code):
        raise HTTPException(400,"规则编号须为 3-32 位大写字母、数字、下划线或短横线")
    if not body.name.strip(): raise HTTPException(400,"规则名称不能为空")
    if body.severity not in RULE_SEVERITIES: raise HTTPException(400,"风险等级无效")
    if body.kind and body.kind not in RULE_KINDS: raise HTTPException(400,"规则类型无效")
    if body.kind=="确定性规则":
        terms=[x.strip() for x in re.split(r"[、,，]",body.fields) if x.strip()]
        if len(terms)<2 or not any(word in body.description for word in ("早于","晚于","先于","不晚于","不得晚于")):
            raise HTTPException(400,"用户确定性规则目前支持日期顺序算法：核验字段至少填写两个日期事件，判定说明需明确‘早于/晚于/先于/不晚于’关系")
    return code

@app.post("/api/rules",status_code=201)
def create_rule(body: RuleWrite):
    code=validate_rule(body,True); stamp=datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    with db() as conn:
        if conn.execute("SELECT 1 FROM rule_definitions WHERE code=?",(code,)).fetchone(): raise HTTPException(409,"规则编号已存在")
        conn.execute("""INSERT INTO rule_definitions(code,name,kind,fields,description,severity,enabled,system_managed,created_at,updated_at)
          VALUES(?,?,?,?,?,?,?,0,?,?)""",(code,body.name.strip(),body.kind,body.fields.strip(),body.description.strip(),body.severity,int(body.enabled),stamp,stamp))
    return {"ok":True,"code":code}

@app.patch("/api/rules/{code}")
def update_rule(code: str, body: RuleWrite):
    validate_rule(body); stamp=datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    with db() as conn:
        current=conn.execute("SELECT * FROM rule_definitions WHERE code=?",(code.upper(),)).fetchone()
        if not current: raise HTTPException(404,"规则不存在")
        kind=current["kind"] if current["system_managed"] else body.kind
        conn.execute("""UPDATE rule_definitions SET name=?,kind=?,fields=?,description=?,severity=?,enabled=?,updated_at=? WHERE code=?""",
                     (body.name.strip(),kind,body.fields.strip(),body.description.strip(),body.severity,int(body.enabled),stamp,code.upper()))
    return {"ok":True,"code":code.upper()}

@app.delete("/api/rules/{code}")
def delete_rule(code: str):
    with db() as conn:
        if not conn.execute("SELECT 1 FROM rule_definitions WHERE code=?",(code.upper(),)).fetchone(): raise HTTPException(404,"规则不存在")
        conn.execute("DELETE FROM rule_definitions WHERE code=?",(code.upper(),))
    return {"ok":True,"code":code.upper()}

@app.get("/api/policies")
def policies():
    with db() as conn:
        items=rows(conn.execute("SELECT id,title,clause,text,source,effective_date,is_template,enabled FROM policy_chunks ORDER BY id").fetchall())
    return {"items":items,"total":len(items)}

class PolicyWrite(BaseModel):
    title: str
    clause: str
    text: str
    source: str
    effective_date: str | None = None
    is_template: bool = False
    enabled: bool = True

def validate_policy(body: PolicyWrite):
    if not all((body.title.strip(),body.clause.strip(),body.text.strip(),body.source.strip())):
        raise HTTPException(400,"政策标题、条款分类、正文和来源不能为空")

@app.post("/api/policies",status_code=201)
def create_policy(body: PolicyWrite):
    validate_policy(body)
    with db() as conn:
        cur=conn.execute("""INSERT INTO policy_chunks(title,clause,text,source,effective_date,is_template,enabled)
          VALUES(?,?,?,?,?,?,?)""",(body.title.strip(),body.clause.strip(),body.text.strip(),body.source.strip(),body.effective_date,int(body.is_template),int(body.enabled)))
    return {"ok":True,"id":cur.lastrowid}

@app.patch("/api/policies/{policy_id}")
def update_policy(policy_id: int, body: PolicyWrite):
    validate_policy(body)
    with db() as conn:
        if not conn.execute("SELECT 1 FROM policy_chunks WHERE id=?",(policy_id,)).fetchone(): raise HTTPException(404,"政策条款不存在")
        conn.execute("""UPDATE policy_chunks SET title=?,clause=?,text=?,source=?,effective_date=?,is_template=?,enabled=? WHERE id=?""",
                     (body.title.strip(),body.clause.strip(),body.text.strip(),body.source.strip(),body.effective_date,int(body.is_template),int(body.enabled),policy_id))
    return {"ok":True,"id":policy_id}

@app.delete("/api/policies/{policy_id}")
def delete_policy(policy_id: int):
    with db() as conn:
        if not conn.execute("SELECT 1 FROM policy_chunks WHERE id=?",(policy_id,)).fetchone(): raise HTTPException(404,"政策条款不存在")
        conn.execute("DELETE FROM policy_chunks WHERE id=?",(policy_id,))
    return {"ok":True,"id":policy_id}

@app.get("/api/projects")
def projects(q: str=""):
    sql="SELECT id,name,category,community,budget,contract_amount,status,progress,risk_count,high_risk_count,updated_at,is_demo FROM projects"
    args=[]
    if q: sql+=" WHERE name LIKE ? OR id LIKE ? OR community LIKE ?"; args=[f"%{q}%"]*3
    sql+=" ORDER BY high_risk_count DESC,risk_count DESC,updated_at DESC"
    with db() as conn: items=rows(conn.execute(sql,args).fetchall())
    return {"items":items,"total":len(items)}

class ProjectCreate(BaseModel):
    id: str
    name: str
    category: str = "公共设施"
    community: str
    budget: int
    contract_amount: int = 0
    status: str = "立项审批"
    progress: int = 0
    with_demo_materials: bool = False
    is_demo: bool = False

class ProjectUpdate(BaseModel):
    name: str | None = None
    category: str | None = None
    community: str | None = None
    budget: int | None = None
    contract_amount: int | None = None
    status: str | None = None
    progress: int | None = None

@app.post("/api/projects",status_code=201)
def create_project(body: ProjectCreate):
    project_id=body.id.strip().upper(); name=body.name.strip(); community=body.community.strip()
    if not project_id or not name or not community: raise HTTPException(400,"项目编号、名称和社区不能为空")
    if body.budget<=0 or body.contract_amount<0: raise HTTPException(400,"金额必须为有效非负数")
    if not 0<=body.progress<=100: raise HTTPException(400,"进度必须在0到100之间")
    contract=body.contract_amount
    facts={}
    project_tuple=(project_id,name,body.category,community,body.budget,contract,body.status,body.progress,0,0)
    with db() as conn:
        if conn.execute("SELECT 1 FROM projects WHERE id=?",(project_id,)).fetchone(): raise HTTPException(409,"项目编号已存在")
        now=datetime.now().strftime("%Y-%m-%d %H:%M")
        conn.execute("""INSERT INTO projects(id,name,category,community,budget,contract_amount,status,progress,risk_count,high_risk_count,updated_at,facts_json,is_demo)
          VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?)""",(*project_tuple,now,json.dumps(facts,ensure_ascii=False),1 if (body.is_demo or body.with_demo_materials) else 0))
        if body.with_demo_materials:
            for filename,doc_type,source in DOC_TYPES:
                doc_lines=lines_for(project_tuple,doc_type); path=FILES_DIR/project_id/filename; make_pdf(path,filename[:-4],doc_lines)
                sha=hashlib.sha256(path.read_bytes()).hexdigest()
                content={"pages":[{"page":1,"width":595,"height":842,"lines":[{"no":i+1,"text":line,"bbox":[67,90+i*28,520,110+i*28]} for i,line in enumerate(doc_lines)]}],"extractor":"generated-demo-v1","language":"zh-CN"}
                conn.execute("INSERT INTO documents(project_id,name,doc_type,source_system,pages,status,file_path,content_json,sha256,created_at) VALUES(?,?,?,?,?,?,?,?,?,?)",(project_id,filename,doc_type,source,1,"已解析",str(path),json.dumps(content,ensure_ascii=False),sha,now))
    return {"ok":True,"project_id":project_id,"documents_created":9 if body.with_demo_materials else 0}

@app.patch("/api/projects/{project_id}")
def update_project(project_id: str, body: ProjectUpdate):
    values=body.model_dump(exclude_none=True)
    if not values: raise HTTPException(400,"没有可更新字段")
    if values.get("budget",1)<=0 or values.get("contract_amount",0)<0: raise HTTPException(400,"金额必须为有效非负数")
    if "progress" in values and not 0<=values["progress"]<=100: raise HTTPException(400,"进度必须在0到100之间")
    for key in ("name","category","community","status"):
        if key in values:
            values[key]=values[key].strip()
            if not values[key]: raise HTTPException(400,f"{key} 不能为空")
    assignments=",".join(f"{key}=?" for key in values)
    with db() as conn:
        if not conn.execute("SELECT 1 FROM projects WHERE id=?",(project_id,)).fetchone(): raise HTTPException(404,"项目不存在")
        conn.execute(f"UPDATE projects SET {assignments},updated_at=? WHERE id=?",(*values.values(),datetime.now().strftime("%Y-%m-%d %H:%M"),project_id))
    return {"ok":True,"project_id":project_id,"updated_fields":list(values)}

@app.delete("/api/projects/{project_id}")
def delete_project(project_id: str, confirm: str=""):
    if confirm!=project_id: raise HTTPException(400,"删除确认值必须与项目编号完全一致")
    with db() as conn:
        row=conn.execute("SELECT name FROM projects WHERE id=?",(project_id,)).fetchone()
        if not row: raise HTTPException(404,"项目不存在")
        run_ids=[r[0] for r in conn.execute("SELECT id FROM audit_runs WHERE project_id=?",(project_id,))]
        trace_paths=[]
        for run_id in run_ids:
            trace_paths.extend(Path(r[0]) for r in conn.execute("SELECT file_path FROM trace_artifacts WHERE run_id=? AND file_path IS NOT NULL",(run_id,)))
            for table in ("audit_events","ai_calls","extracted_facts"):
                conn.execute(f"DELETE FROM {table} WHERE run_id=?",(run_id,))
            conn.execute("UPDATE run_spans SET input_artifact_id=NULL,output_artifact_id=NULL,error_artifact_id=NULL WHERE run_id=?",(run_id,))
            conn.execute("DELETE FROM trace_artifacts WHERE run_id=?",(run_id,))
            conn.execute("DELETE FROM run_spans WHERE run_id=?",(run_id,))
            conn.execute("DELETE FROM audit_run_documents WHERE run_id=?",(run_id,))
        conn.execute("DELETE FROM risk_transitions WHERE project_id=?",(project_id,))
        # Children reference their immutable parent versions. Delete the newest
        # descendants first so a project with re-normalized/MinerU comparison
        # versions can be removed without violating the self foreign key.
        version_ids=[r[0] for r in conn.execute("SELECT v.id FROM document_parser_versions v JOIN documents d ON d.id=v.document_id WHERE d.project_id=? ORDER BY v.id DESC",(project_id,))]
        conn.execute("UPDATE documents SET active_parser_version_id=NULL WHERE project_id=?",(project_id,))
        for version_id in version_ids:
            for table in ("visual_element_analyses","document_chunks","document_elements","document_parse_jobs"):
                conn.execute(f"DELETE FROM {table} WHERE parser_version_id=?",(version_id,))
            conn.execute("DELETE FROM document_parser_versions WHERE id=?",(version_id,))
        for table in ("anomalies","audit_runs","review_actions","timeline","documents"):
            conn.execute(f"DELETE FROM {table} WHERE project_id=?",(project_id,))
        conn.execute("DELETE FROM projects WHERE id=?",(project_id,))
    for path in trace_paths:
        if path.exists(): path.unlink()
    source=FILES_DIR/project_id; archived=None
    if source.exists():
        trash=FILES_DIR.parent/"trash"; trash.mkdir(parents=True,exist_ok=True)
        archived=trash/f"{project_id}-{datetime.now().strftime('%Y%m%d-%H%M%S')}"; shutil.move(str(source),str(archived))
    return {"ok":True,"project_id":project_id,"project_name":row["name"],"database_deleted_permanently":True,
            "trace_runs_deleted":len(run_ids),"parser_versions_deleted":len(version_ids),
            "original_files_recoverable":bool(archived),"files_archived":str(archived) if archived else None}

@app.get("/api/projects/{project_id}")
def project_detail(project_id: str):
    with db() as conn:
        if not conn.execute("SELECT 1 FROM projects WHERE id=?",(project_id,)).fetchone(): raise HTTPException(404,"项目不存在")
        legacy_ids=[row["id"] for row in conn.execute("SELECT id FROM documents WHERE project_id=? AND deleted_at IS NULL AND active_parser_version_id IS NULL",(project_id,))]
    for document_id in legacy_ids:
        try: ensure_legacy_version(document_id)
        except Exception: pass
    with db() as conn:
        p=conn.execute("SELECT * FROM projects WHERE id=?",(project_id,)).fetchone()
        if not p: raise HTTPException(404,"项目不存在")
        project=decode(p,"facts_json"); facts=project.pop("facts_json")
        latest_run=conn.execute("SELECT id FROM audit_runs WHERE project_id=? AND status='completed' AND run_kind='audit' ORDER BY id DESC LIMIT 1",(project_id,)).fetchone()
        anomalies=[]
        anomaly_rows=conn.execute("SELECT * FROM anomalies WHERE run_id=? ORDER BY CASE severity WHEN 'high' THEN 0 WHEN 'moderate' THEN 1 WHEN 'minor' THEN 2 WHEN 'medium' THEN 3 ELSE 4 END,id",(latest_run["id"] if latest_run else -1,))
        for row in anomaly_rows:
            a=decode(row,"evidence_json"); a["evidences"]=a.pop("evidence_json"); anomalies.append(a)
        documents=rows(conn.execute("""SELECT d.id,d.name,d.doc_type,d.pages,d.source_system,d.status,d.sha256,d.active_parser_version_id,
          v.parser_kind AS active_parser_kind,v.parser_name AS active_parser_name,v.parser_version AS active_parser_version,v.provider_id AS active_parser_provider_id,v.status AS parser_status,
          v.stats_json AS parser_stats_json,v.warnings_json AS parser_warnings_json,v.trace_run_id AS parser_trace_run_id,
          (SELECT COUNT(*) FROM document_parser_versions x WHERE x.document_id=d.id) AS parser_version_count
          FROM documents d LEFT JOIN document_parser_versions v ON v.id=d.active_parser_version_id WHERE d.project_id=? AND d.deleted_at IS NULL ORDER BY d.id""",(project_id,)).fetchall())
        for document in documents:
            document["parser_stats"]=json.loads(document.pop("parser_stats_json") or "{}")
            document["parser_warnings"]=json.loads(document.pop("parser_warnings_json") or "[]")
            frozen=conn.execute("SELECT parser_version_id FROM audit_run_documents WHERE run_id=? AND document_id=?",(latest_run["id"] if latest_run else -1,document["id"])).fetchone()
            document["last_audit_parser_version_id"]=frozen["parser_version_id"] if frozen else None
            document["parser_stale"]=bool(frozen and frozen["parser_version_id"]!=document["active_parser_version_id"])
        timeline=[{"date":r["event_date"],"title":r["title"],"source":r["source"],"state":r["state"]} for r in conn.execute("SELECT * FROM timeline WHERE project_id=? ORDER BY event_date",(project_id,))]
    contract=project["contract_amount"]
    paid=float(facts.get("paid_amount",0) or 0); change=float(facts.get("change_amount",0) or 0); ratio=float(facts.get("progress_payment_ratio",0) or 0)
    metrics={"paid_amount":paid,"payment_rate":round(paid/contract,4) if contract else 0,"change_amount":change,"change_rate":round(change/contract,4) if contract else 0,"allowed_payment":round(contract*ratio)}
    return {"project":project,"metrics":metrics,"facts":facts,"anomalies":anomalies,"documents":documents,"timeline":timeline}

class AuditStart(BaseModel):
    route_overrides: dict | None = None
    parent_run_id: int | None = None
    derived_from_stage: str | None = None
    run_kind: str = "audit"
    parser_versions: dict[str, int] | None = None
    review_mode: str = "final"
    action_type: str | None = None
    action_title: str | None = None
    planned_at: str | None = None
    cutoff_at: str | None = None
    current_document_ids: list[int] | None = None


@app.get("/api/review-actions/catalog")
def review_action_catalog():
    return {"modes": [
        {"value": "gate", "label": "阶段门禁审查", "description": "在付款、变更、验收等动作批准前给出放行建议"},
        {"value": "incremental", "label": "新增资料增量审查", "description": "只把本批资料作为变化输入，并结合截止时点的历史证据"},
        {"value": "final", "label": "全过程全量终审", "description": "竣工结算、年度内控或历史项目审计"},
    ], "actions": [{"value": key, **value} for key, value in ACTION_TYPES.items()]}


@app.get("/api/projects/{project_id}/review-actions")
def project_review_actions(project_id: str):
    with db() as conn:
        if not conn.execute("SELECT 1 FROM projects WHERE id=?", (project_id,)).fetchone():
            raise HTTPException(404, "项目不存在")
        items=[]
        for row in conn.execute("""SELECT a.*,r.id AS run_id,r.status AS run_status,r.decision_json
          FROM review_actions a LEFT JOIN audit_runs r ON r.action_id=a.id
          WHERE a.project_id=? ORDER BY a.id DESC""", (project_id,)):
            item=decode(row,"current_document_ids_json","decision_json")
            items.append(item)
    return {"items":items,"total":len(items)}


@app.post("/api/projects/{project_id}/audit",status_code=202)
def run_audit(project_id: str, background_tasks: BackgroundTasks, body: AuditStart | None = None):
    body = body or AuditStart()
    body.route_overrides = body.route_overrides or {}
    try:
        review_mode,action_type=normalized_action(body.review_mode,body.action_type)
        cutoff_at=parse_cutoff(body.cutoff_at)
    except ValueError as exc:
        raise HTTPException(400,str(exc))
    current_document_ids={int(value) for value in (body.current_document_ids or [])}
    if review_mode in {"gate","incremental"} and not current_document_ids:
        raise HTTPException(400,"阶段门禁或增量审查至少选择一份本批资料")
    invalid = set(body.route_overrides) - set(ALL_STAGES)
    if invalid: raise HTTPException(400,f"未知阶段：{', '.join(sorted(invalid))}")
    config_snapshot = snapshot(body.route_overrides)
    prompts = prompt_snapshot()
    route_label = ",".join(sorted({f"{v.get('provider_id')}:{v.get('model')}" for v in config_snapshot["stage_routes"].values()}))
    with db() as conn:
        document_ids=[row["id"] for row in conn.execute("SELECT id FROM documents WHERE project_id=? AND deleted_at IS NULL AND created_at<=?",(project_id,cutoff_at))]
    invalid_current=current_document_ids-set(document_ids)
    if invalid_current: raise HTTPException(400,"本批资料包含不属于当前项目或晚于截止时间的文档")
    for document_id in document_ids:
        try: ensure_legacy_version(document_id)
        except Exception: pass
    with db() as conn:
        if not conn.execute("SELECT 1 FROM projects WHERE id=?",(project_id,)).fetchone(): raise HTTPException(404,"项目不存在")
        running=conn.execute("SELECT id FROM audit_runs WHERE project_id=? AND run_kind='audit' AND status IN ('queued','running') ORDER BY id DESC LIMIT 1",(project_id,)).fetchone()
        if running: return {"ok":True,"run_id":running["id"],"status":"running"}
        baseline=conn.execute("SELECT id FROM audit_runs WHERE project_id=? AND status='completed' AND run_kind='audit' ORDER BY id DESC LIMIT 1",(project_id,)).fetchone()
        selected_versions=[];overrides=body.parser_versions or {}
        for document in conn.execute("SELECT id,name,active_parser_version_id FROM documents WHERE project_id=? AND deleted_at IS NULL AND created_at<=? ORDER BY id",(project_id,cutoff_at)):
            version_id=int(overrides.get(str(document["id"])) or document["active_parser_version_id"] or 0)
            version=conn.execute("SELECT id,status FROM document_parser_versions WHERE id=? AND document_id=?",(version_id,document["id"])).fetchone()
            if not version or version["status"] not in {"ready","ready_with_warnings"}:
                raise HTTPException(409,f"资料“{document['name']}”尚未选择可用的 active 解析版本")
            selected_versions.append((document["id"],version_id,1 if str(document["id"]) in overrides else 0,
                                      "current" if document["id"] in current_document_ids else "history"))
        if not selected_versions: raise HTTPException(409,"项目没有可审查资料")
        started=datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        definition=action_definition(action_type)
        title=(body.action_title or definition["label"]).strip()
        action_snapshot={"review_mode":review_mode,"action_type":action_type,"action_label":definition["label"],
                         "title":title,"planned_at":body.planned_at,"cutoff_at":cutoff_at,
                         "required_phases":definition["required_phases"],"current_document_ids":sorted(current_document_ids)}
        action_cur=conn.execute("""INSERT INTO review_actions(project_id,action_type,title,planned_at,cutoff_at,status,current_document_ids_json,created_at,updated_at)
          VALUES(?,?,?,?,?,'under_review',?,?,?)""",(project_id,action_type,title,body.planned_at,cutoff_at,
          json.dumps(sorted(current_document_ids)),started,started))
        action_id=action_cur.lastrowid
        cur=conn.execute("""INSERT INTO audit_runs(project_id,started_at,status,rule_count,fact_count,anomaly_count,provider,progress,current_stage,result_json,
          parent_run_id,derived_from_stage,run_kind,config_snapshot_json,prompt_versions_json,route_overrides_json,
          action_id,review_mode,cutoff_at,baseline_run_id,action_snapshot_json,decision_json)
          VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",(project_id,started,"running",0,0,0,route_label,0,"queued","{}",
          body.parent_run_id,body.derived_from_stage,body.run_kind,json.dumps(config_snapshot,ensure_ascii=False),
          json.dumps(prompts,ensure_ascii=False),json.dumps(body.route_overrides,ensure_ascii=False),action_id,review_mode,cutoff_at,
          baseline["id"] if baseline else None,json.dumps(action_snapshot,ensure_ascii=False),"{}"))
        run_id=cur.lastrowid
        conn.executemany("INSERT INTO audit_run_documents(run_id,document_id,parser_version_id,is_override,document_role) VALUES(?,?,?,?,?)",[(run_id,*item) for item in selected_versions])
        conn.execute("INSERT INTO audit_events(run_id,sequence,created_at,stage,status,message,detail_json) VALUES(?,?,?,?,?,?,?)",
                     (run_id,1,started,"queued","completed",f"{definition['label']}任务已创建，资料截止 {cutoff_at}",json.dumps(action_snapshot,ensure_ascii=False)))
    background_tasks.add_task(execute_audit,run_id,project_id)
    return {"ok":True,"run_id":run_id,"status":"running","poll_url":f"/api/audit-runs/{run_id}"}

def audit_run_payload(conn, run_id):
    run=conn.execute("SELECT * FROM audit_runs WHERE id=?",(run_id,)).fetchone()
    if not run: raise HTTPException(404,"审查运行不存在")
    item=decode(run,"result_json","config_snapshot_json","prompt_versions_json","route_overrides_json",
                "action_snapshot_json","decision_json")
    events=[decode(row,"detail_json") for row in conn.execute("SELECT * FROM audit_events WHERE run_id=? ORDER BY sequence",(run_id,))]
    calls=rows(conn.execute("SELECT id,stage,provider,model,started_at,duration_ms,success,input_tokens,output_tokens,request_hash,response_preview,error FROM ai_calls WHERE run_id=? ORDER BY id",(run_id,)).fetchall())
    spans=[decode(row,"metadata_json") for row in conn.execute("SELECT * FROM run_spans WHERE run_id=? ORDER BY sequence",(run_id,))]
    facts=[]
    for row in conn.execute("SELECT f.*,d.name AS document_name,COALESCE(f.document_phase,d.doc_type,'其他资料') AS document_type FROM extracted_facts f LEFT JOIN documents d ON d.id=f.document_id WHERE f.run_id=? ORDER BY f.id",(run_id,)):
        facts.append(decode(row,"value_json","element_ids_json"))
    anomalies=[]
    for row in conn.execute("SELECT * FROM anomalies WHERE run_id=? ORDER BY CASE severity WHEN 'high' THEN 0 WHEN 'moderate' THEN 1 WHEN 'minor' THEN 2 WHEN 'medium' THEN 3 ELSE 4 END,id",(run_id,)):
        a=decode(row,"evidence_json"); a["evidences"]=a.pop("evidence_json"); anomalies.append(a)
    transitions=rows(conn.execute("SELECT * FROM risk_transitions WHERE run_id=? ORDER BY id",(run_id,)).fetchall())
    return {"run":item,"events":events,"calls":calls,"spans":spans,"facts":facts,"anomalies":anomalies,
            "risk_transitions":transitions}

@app.get("/api/audit-runs/{run_id}")
def audit_run(run_id: int):
    with db() as conn: return audit_run_payload(conn,run_id)

@app.get("/api/projects/{project_id}/audit-runs/latest")
def latest_audit_run(project_id: str):
    with db() as conn:
        row=conn.execute("SELECT id FROM audit_runs WHERE project_id=? AND run_kind='audit' ORDER BY id DESC LIMIT 1",(project_id,)).fetchone()
        return audit_run_payload(conn,row["id"]) if row else {"run":None,"events":[],"calls":[],"facts":[],"anomalies":[]}

@app.get("/api/anomalies")
def all_anomalies(severity: str="", status: str="", include_demo: bool=False):
    sql="""SELECT a.*,p.name AS project_name,p.community FROM anomalies a
             JOIN projects p ON p.id=a.project_id WHERE a.run_id=(SELECT r.id FROM audit_runs r WHERE r.project_id=a.project_id AND r.status='completed' AND r.run_kind='audit' ORDER BY r.id DESC LIMIT 1)"""
    args=[]
    if not include_demo: sql+=" AND p.is_demo=0"
    if severity: sql+=" AND a.severity=?"; args.append(severity)
    if status: sql+=" AND a.status=?"; args.append(status)
    sql+=" ORDER BY CASE a.severity WHEN 'high' THEN 0 WHEN 'moderate' THEN 1 WHEN 'minor' THEN 2 WHEN 'medium' THEN 3 ELSE 4 END,a.id"
    with db() as conn:
        items=[]
        for row in conn.execute(sql,args):
            item=decode(row,"evidence_json"); item["evidences"]=item.pop("evidence_json"); items.append(item)
    return {"items":items,"total":len(items)}

class AnomalyUpdate(BaseModel):
    status: str

@app.patch("/api/anomalies/{anomaly_id}")
def update_anomaly(anomaly_id: int, body: AnomalyUpdate):
    allowed={"待复核","已确认","已转交","误报"}
    if body.status not in allowed: raise HTTPException(400,"不支持的处理状态")
    with db() as conn:
        anomaly=conn.execute("SELECT project_id,run_id FROM anomalies WHERE id=?",(anomaly_id,)).fetchone()
        if not anomaly: raise HTTPException(404,"异常不存在")
        cur=conn.execute("UPDATE anomalies SET status=? WHERE id=?",(body.status,anomaly_id))
        if not cur.rowcount: raise HTTPException(404,"异常不存在")
        latest=conn.execute("SELECT id FROM audit_runs WHERE project_id=? AND status='completed' AND run_kind='audit' ORDER BY id DESC LIMIT 1",(anomaly["project_id"],)).fetchone()
        if latest:
            pending=conn.execute("SELECT COUNT(*) FROM anomalies WHERE run_id=? AND status='待复核'",(latest["id"],)).fetchone()[0]
            high=conn.execute("SELECT COUNT(*) FROM anomalies WHERE run_id=? AND status='待复核' AND severity='high'",(latest["id"],)).fetchone()[0]
            conn.execute("UPDATE projects SET risk_count=?,high_risk_count=?,updated_at=? WHERE id=?",
                         (pending,high,datetime.now().strftime("%Y-%m-%d %H:%M"),anomaly["project_id"]))
    return {"ok":True,"id":anomaly_id,"status":body.status}

@app.get("/api/projects/{project_id}/report")
def audit_report(project_id: str):
    with db() as conn:
        p=conn.execute("SELECT * FROM projects WHERE id=?",(project_id,)).fetchone()
        if not p: raise HTTPException(404,"项目不存在")
        run=conn.execute("SELECT * FROM audit_runs WHERE project_id=? AND run_kind='audit' ORDER BY id DESC LIMIT 1",(project_id,)).fetchone()
        anomalies=[]
        latest_run_id=run["id"] if run else -1
        for row in conn.execute("SELECT * FROM anomalies WHERE run_id=? ORDER BY id",(latest_run_id,)):
            item=decode(row,"evidence_json"); anomalies.append(item)
        calls=conn.execute("SELECT provider,model,stage,duration_ms,success FROM ai_calls WHERE run_id=? ORDER BY id",(run["id"],)).fetchall() if run else []
        documents=conn.execute("SELECT name,doc_type,pages,sha256 FROM documents WHERE project_id=? AND deleted_at IS NULL ORDER BY id",(project_id,)).fetchall()
    buf=io.BytesIO(); c=canvas.Canvas(buf,pagesize=A4); width,height=A4
    try: pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light")); font="STSong-Light"
    except Exception: font="Helvetica"
    def write(text,size=9,indent=0,gap=17):
        nonlocal y
        value=str(text or ""); chunks=[value[i:i+52] for i in range(0,len(value),52)] or [""]
        for line in chunks:
            if y<65: c.showPage(); c.setFont(font,size); y=height-60
            c.setFont(font,size); c.drawString(58+indent,y,line); y-=gap
    c.setTitle(f"{project_id}智能审查报告"); c.setFont(font,18); c.drawCentredString(width/2,height-55,"农村集体建设工程全过程智能审查报告")
    c.setFont(font,9); y=height-90
    for line in [f"项目编号：{p['id']}    项目名称：{p['name']}",f"所属社区：{p['community']}    工程类别：{p['category']}",
                 f"项目建档预算：{p['budget']:,} 元    项目建档合同额：{p['contract_amount']:,} 元",
                 f"最新审查运行：#{run['id']}  {run['status']}    开始：{run['started_at']}    完成：{run['finished_at']}" if run else "尚无审查运行",
                 f"资料：{len(documents)} 份    事实：{run['fact_count'] if run else 0} 条    规则：{run['rule_count'] if run else 0} 条    待复核：{len(anomalies)} 项"]: write(line,9)
    y-=8; write("一、审查技术链路",13,gap=24)
    write("原件版面解析 → 按阶段动态选择 OpenAI-compatible 模型 → Python/SQL 去重、计算与日期规则 → 语义一致性复核 → 政策检索。",9)
    write(f"模型调用共 {len(calls)} 次；"+"；".join(f"{x['stage']} / {x['provider']} / {'成功' if x['success'] else '失败'} / {x['duration_ms']}ms" for x in calls),8)
    y-=8; write("二、异常清单",13,gap=24)
    if not anomalies: write("本次未产生待人工复核事项。",9)
    for i,a in enumerate(anomalies,1):
        severity={"high":"高风险","moderate":"中风险","minor":"低风险","medium":"需关注","low":"提示"}.get(a["severity"],a["severity"])
        write(f"{i}. [{severity}] {a['code']} {a['title']}（{a['status']}，置信度 {a['confidence']:.0%}）",10)
        write(a["summary"],9,12)
        for ev in a.get("evidence_json") or []:
            write(f"证据：{ev.get('document_name','')} P{ev.get('page',1)}"+(f" B{ev.get('block_id')}" if ev.get('block_id') else "")+f"｜{ev.get('quote','')}",8,20,15)
        y-=5
    y-=5; write("三、资料清单",13,gap=24)
    for i,d in enumerate(documents,1): write(f"{i}. [{d['doc_type']}] {d['name']}｜{d['pages']}页｜SHA256 {d['sha256'][:16]}…",8)
    c.setFont(font,8); c.setFillColorRGB(.4,.45,.5); c.drawString(58,35,"说明：本报告记录机器审查结果与证据来源，不替代工作人员最终判断；政策模板不作为正式制度依据。")
    c.save(); data=buf.getvalue()
    return Response(data,media_type="application/pdf",headers={"Content-Disposition":f'attachment; filename="audit-report-{project_id}.pdf"'})

@app.get("/api/documents/{document_id}")
def document(document_id: int):
    with db() as conn:
        row=conn.execute("SELECT * FROM documents WHERE id=?",(document_id,)).fetchone()
        if not row: raise HTTPException(404,"资料不存在")
        result=decode(row,"content_json"); result["content"]=result.pop("content_json"); result["download_url"]=f"/api/documents/{document_id}/download"; return result

@app.get("/api/documents/{document_id}/download")
def download(document_id: int):
    with db() as conn: row=conn.execute("SELECT name,file_path FROM documents WHERE id=?",(document_id,)).fetchone()
    if not row or not row["file_path"] or not Path(row["file_path"]).exists(): raise HTTPException(404,"原件不存在")
    return FileResponse(row["file_path"],media_type="application/pdf",filename=row["name"])

def extract_local(path: Path):
    return parse_document(path)

def infer_document_type(filename: str, requested: str):
    if requested in CANONICAL_PHASES:
        return requested
    if requested and requested not in ("自动识别","其他资料"):
        probe={"name":filename,"doc_type":requested}
        inferred=canonical_phase(probe)
        return requested if inferred=="其他资料" else inferred
    return canonical_phase({"name":filename,"doc_type":""})

def store_upload(project_id: str, doc_type: str, source_system: str, file: UploadFile):
    allowed={".pdf",".docx",".txt",".png",".jpg",".jpeg",".webp",".bmp"}; suffix=Path(file.filename or "").suffix.lower()
    if suffix not in allowed: raise HTTPException(400,f"{file.filename}：仅支持 PDF、DOCX、TXT 和常见图片")
    target_dir=FILES_DIR/project_id;target_dir.mkdir(parents=True,exist_ok=True);target=target_dir/(file.filename or f"upload{suffix}")
    if target.exists():target=target_dir/f"{target.stem}-{datetime.now().strftime('%Y%m%d%H%M%S%f')}{target.suffix}"
    with target.open("wb") as out:shutil.copyfileobj(file.file,out)
    content=target.read_bytes();digest=hashlib.sha256(content).hexdigest();canonical=infer_document_type(target.name,doc_type)
    with db() as conn:
        existing=conn.execute("SELECT * FROM documents WHERE project_id=? AND sha256=? AND deleted_at IS NULL",(project_id,digest)).fetchone()
        if existing:
            target.unlink(missing_ok=True);return dict(existing),True
        cur=conn.execute("""INSERT INTO documents(project_id,name,doc_type,source_system,pages,status,file_path,content_json,sha256,created_at)
          VALUES(?,?,?,?,0,'待解析',?,'{}',?,?)""",(project_id,target.name,canonical,source_system,str(target),digest,datetime.now().strftime("%Y-%m-%d %H:%M:%S")))
        return dict(conn.execute("SELECT * FROM documents WHERE id=?",(cur.lastrowid,)).fetchone()),False


async def create_upload_versions(project_id: str, stored: list[tuple[dict,bool]], parse_mode: str,
                                 background_tasks: BackgroundTasks, force: bool=False,
                                 parser_provider_id: str|None=None, parser_model: str|None=None,
                                 auto_activate: bool=True):
    if parse_mode not in {"pymupdf","mineru"}:raise HTTPException(400,"解析方案必须是 pymupdf 或 mineru")
    items=[]
    if parse_mode=="pymupdf":
        for document,duplicate in stored:
            version=create_version(document["id"],"pymupdf",PYMUPDF_VERSION,"pymupdf",PYMUPDF_VERSION,force=force)
            if version["status"] not in {"ready","ready_with_warnings"}:run_pymupdf_version(version["id"],auto_activate)
            elif auto_activate:activate_version(document["id"],version["id"])
            items.append({"document_id":document["id"],"name":document["name"],"duplicate":duplicate,"parser_version_id":version["id"],"parser_status":"ready","active":auto_activate})
        return items
    overrides={"ocr":{"provider_id":parser_provider_id,"model":parser_model}} if parser_provider_id else None
    provider,model=resolve_route("ocr",overrides)
    if not str(provider.get("kind","")).startswith("mineru"):raise HTTPException(400,"OCR 解析必须选择 MinerU 供应商")
    if provider["kind"]=="mineru-local":
        model=parser_model or "vlm-engine"
        if model!="vlm-engine":raise HTTPException(400,"本地 MinerU 仅允许 vlm-engine；hybrid-engine 已禁用以避免 OOM")
    version_rows=[]
    for document,duplicate in stored:
        version=create_version(document["id"],"mineru",f"mineru-{model}-{datetime.now().strftime('%Y%m%d')}",provider["id"],model,{"enable_table":True,"enable_formula":True,"language":"ch","auto_activate":auto_activate},force=force)
        if version["status"] in {"ready","ready_with_warnings"}:
            if auto_activate:activate_version(document["id"],version["id"])
            items.append({"document_id":document["id"],"name":document["name"],"duplicate":duplicate,"parser_version_id":version["id"],"parser_status":version["status"],"active":auto_activate or bool(version["is_active"]),"reused":True});continue
        data_id=f"{project_id}-{document['id']}-{document['sha256'][:12]}"
        with db() as conn:
            attempt=conn.execute("SELECT COALESCE(MAX(attempt),0)+1 FROM document_parse_jobs WHERE parser_version_id=?",(version["id"],)).fetchone()[0]
            conn.execute("""INSERT INTO document_parse_jobs(parser_version_id,attempt,provider_id,data_id,status,progress,created_at,updated_at)
              VALUES(?,?,?,?,?,?,?,?)""",(version["id"],attempt,provider["id"],data_id,"uploading",10,datetime.now().strftime("%Y-%m-%d %H:%M:%S"),datetime.now().strftime("%Y-%m-%d %H:%M:%S")))
        version_rows.append((document,version,data_id,duplicate))
    if version_rows:
        try:
            submitted=await mineru_upload_files([(Path(document["file_path"]),data_id) for document,version,data_id,duplicate in version_rows],model,8,provider["id"])
        except Exception as exc:
            with db() as conn:
                for document,version,data_id,duplicate in version_rows:
                    conn.execute("UPDATE document_parser_versions SET status='failed',warnings_json=? WHERE id=?",(json.dumps([str(exc)],ensure_ascii=False),version["id"]));conn.execute("UPDATE document_parse_jobs SET status='upload_failed',error_json=?,updated_at=?,finished_at=? WHERE parser_version_id=?",(json.dumps({"type":type(exc).__name__,"message":str(exc)},ensure_ascii=False),datetime.now().strftime("%Y-%m-%d %H:%M:%S"),datetime.now().strftime("%Y-%m-%d %H:%M:%S"),version["id"]))
            raise HTTPException(502,f"MinerU 批量提交失败：{exc}")
        versions_by_name={};trace_runs=[]
        with db() as conn:
            for document,version,data_id,duplicate in version_rows:
                versions_by_name[document["name"]]=version["id"]
                conn.execute("UPDATE document_parser_versions SET status='submitted' WHERE id=?",(version["id"],));conn.execute("UPDATE document_parse_jobs SET batch_id=?,status='submitted',progress=25,request_json=?,response_json=?,updated_at=? WHERE parser_version_id=?",(submitted["batch_id"],json.dumps(submitted["request"],ensure_ascii=False),json.dumps(submitted["response"],ensure_ascii=False),datetime.now().strftime("%Y-%m-%d %H:%M:%S"),version["id"]))
                trace_runs.append(version["trace_run_id"])
                items.append({"document_id":document["id"],"name":document["name"],"duplicate":duplicate,"parser_version_id":version["id"],"parser_status":"submitted","active":False,"batch_id":submitted["batch_id"]})
        for trace_run_id in trace_runs:
            span=start_span(trace_run_id,"mineru_submit","provider","MinerU 批量提交",provider_id=provider["id"],model=model);attach_input(span,submitted["request"]);finish_span(span,output=submitted["response"])
        if provider["kind"]=="mineru-local":background_tasks.add_task(poll_mineru_local_task,submitted["batch_id"],versions_by_name,provider["id"])
        else:background_tasks.add_task(poll_mineru_batch,submitted["batch_id"],versions_by_name,provider["id"])
    return items


@app.post("/api/documents/upload")
async def upload(background_tasks: BackgroundTasks,project_id: str=Form(...),doc_type: str=Form("自动识别"),source_system: str=Form("人工上传"),parse_mode: str=Form("pymupdf"),parser_provider_id: str|None=Form(None),parser_model: str|None=Form(None),auto_activate: bool=Form(True),file: UploadFile=File(...)):
    with db() as conn:
        if not conn.execute("SELECT 1 FROM projects WHERE id=?",(project_id,)).fetchone():raise HTTPException(404,"项目不存在")
    items=await create_upload_versions(project_id,[store_upload(project_id,doc_type,source_system,file)],parse_mode,background_tasks,False,parser_provider_id,parser_model,auto_activate)
    return {"ok":True,**items[0]}


@app.post("/api/documents/upload-batch")
async def upload_batch(background_tasks: BackgroundTasks,project_id: str=Form(...),doc_type: str=Form("自动识别"),source_system: str=Form("人工上传"),parse_mode: str=Form("pymupdf"),parser_provider_id: str|None=Form(None),parser_model: str|None=Form(None),auto_activate: bool=Form(True),files: list[UploadFile]=File(...)):
    if not 1<=len(files)<=200:raise HTTPException(400,"单批文件数必须为 1-200")
    with db() as conn:
        if not conn.execute("SELECT 1 FROM projects WHERE id=?",(project_id,)).fetchone():raise HTTPException(404,"项目不存在")
    stored=[]
    try:
        for file in files:stored.append(store_upload(project_id,doc_type,source_system,file))
        items=await create_upload_versions(project_id,stored,parse_mode,background_tasks,False,parser_provider_id,parser_model,auto_activate)
        return {"ok":True,"items":items,"total":len(items),"parse_mode":parse_mode}
    except Exception:
        raise


class ParserVersionCreate(BaseModel):
    parser_kind: str
    force: bool = False
    provider_id: str | None = None
    model: str | None = None
    auto_activate: bool = True


class ParserVersionActivate(BaseModel):
    parser_version_id: int


@app.get("/api/documents/{document_id}/parser-versions")
def parser_versions(document_id: int):
    try: ensure_legacy_version(document_id)
    except Exception: pass
    with db() as conn:
        if not conn.execute("SELECT 1 FROM documents WHERE id=?",(document_id,)).fetchone():raise HTTPException(404,"资料不存在")
        ids=[row["id"] for row in conn.execute("SELECT id FROM document_parser_versions WHERE document_id=? ORDER BY id DESC",(document_id,))]
    return {"items":[version_payload(i,False) for i in ids],"total":len(ids)}


@app.get("/api/parser-versions/{version_id}")
def parser_version_detail(version_id: int, include_elements: bool=True):
    try:return version_payload(version_id,include_elements)
    except KeyError:raise HTTPException(404,"解析版本不存在")


@app.post("/api/documents/{document_id}/parser-versions",status_code=202)
async def create_parser_version(document_id: int, body: ParserVersionCreate, background_tasks: BackgroundTasks):
    with db() as conn:document=conn.execute("SELECT * FROM documents WHERE id=?",(document_id,)).fetchone()
    if not document:raise HTTPException(404,"资料不存在")
    items=await create_upload_versions(document["project_id"],[(dict(document),True)],body.parser_kind,background_tasks,body.force,body.provider_id,body.model,body.auto_activate)
    return {"ok":True,**items[0]}


@app.patch("/api/documents/{document_id}/active-parser-version")
def set_active_parser_version(document_id: int, body: ParserVersionActivate):
    try:activate_version(document_id,body.parser_version_id)
    except KeyError:raise HTTPException(404,"解析版本不存在")
    except ValueError as exc:raise HTTPException(409,str(exc))
    return {"ok":True,"document_id":document_id,"parser_version_id":body.parser_version_id,
            "warning":"现有审查运行仍使用其固化解析版本；请按需重新审查"}


@app.get("/api/documents/{document_id}/parser-compare")
def parser_compare(document_id: int, left: int, right: int):
    a,b=version_payload(left,True),version_payload(right,True)
    if a["document_id"]!=document_id or b["document_id"]!=document_id:raise HTTPException(400,"解析版本不属于当前资料")
    def summary(item):
        elements=item.get("elements") or [];types={}
        for element in elements:types[element["element_type"]]=types.get(element["element_type"],0)+1
        return {"id":item["id"],"parser_name":item["parser_name"],"parser_version":item["parser_version"],"provider_id":item["provider_id"],"model":item["model"],"status":item["status"],
                "is_active":item["is_active"],"stats":item["stats_json"],"warnings":item["warnings_json"],"types":types,
                "text_preview":"\n".join((x.get("markdown") or x.get("text") or "") for x in elements[:30])[:6000]}
    def normalized_text(item):return re.sub(r"\s+","","\n".join((x.get("markdown") or x.get("text") or "") for x in item.get("elements") or []))[:200000]
    def table_cells(item):
        values=set()
        for element in item.get("elements") or []:
            for row in element.get("cell_grid_json") or []:
                for cell in row:
                    value=re.sub(r"\s+","",str(cell.get("text") or ""))
                    if value:values.add(value)
        return values
    left_text,right_text=normalized_text(a),normalized_text(b);left_cells,right_cells=table_cells(a),table_cells(b)
    left_types=set((a.get("stats_json") or {}).get("types") or {});right_types=set((b.get("stats_json") or {}).get("types") or {})
    def jaccard(x,y):return round(len(x&y)/len(x|y),4) if x or y else 1.0
    def coordinate_coverage(item):
        elements=item.get("elements") or [];return round(sum(1 for x in elements if x.get("bbox_json"))/len(elements),4) if elements else 0
    return {"document_id":document_id,"left":summary(a),"right":summary(b),
            "differences":{"element_count":len(a.get("elements") or [])-len(b.get("elements") or []),
                           "coordinate_count":a["stats_json"].get("coordinate_elements",0)-b["stats_json"].get("coordinate_elements",0),
                           "table_count":a["stats_json"].get("tables",0)-b["stats_json"].get("tables",0),
                           "visual_count":a["stats_json"].get("visual_elements",0)-b["stats_json"].get("visual_elements",0)},
            "alignment":{"text_similarity":round(SequenceMatcher(None,left_text,right_text,autojunk=False).ratio(),4) if left_text or right_text else 1.0,
                         "element_type_jaccard":jaccard(left_types,right_types),"table_cell_jaccard":jaccard(left_cells,right_cells),
                         "page_count_match":a["stats_json"].get("pages")==b["stats_json"].get("pages"),
                         "left_coordinate_coverage":coordinate_coverage(a),"right_coordinate_coverage":coordinate_coverage(b),
                         "note":"对齐表示统一数据模型可比，不要求云端与本地版面分块逐元素完全相同"}}


@app.get("/api/parser-versions/{version_id}/raw-zip")
def parser_raw_zip(version_id: int):
    version=version_payload(version_id,False);path=Path(version.get("artifact_dir") or "")/"raw"/"mineru-result.zip"
    if version["parser_kind"]!="mineru" or not path.exists():raise HTTPException(404,"该版本没有 MinerU 原始 ZIP")
    return FileResponse(path,media_type="application/zip",filename=f"parser-version-{version_id}-mineru.zip")


@app.post("/api/parser-versions/{version_id}/renormalize",status_code=202)
def renormalize_parser_version(version_id: int, background_tasks: BackgroundTasks):
    try:child=create_renormalized_version(version_id)
    except KeyError:raise HTTPException(404,"Parser version not found")
    except FileNotFoundError as exc:raise HTTPException(409,str(exc))
    except ValueError as exc:raise HTTPException(409,str(exc))
    background_tasks.add_task(run_renormalized_version,version_id,int(child["id"]))
    return {"ok":True,"source_parser_version_id":version_id,"parser_version_id":child["id"],
            "parent_version_id":version_id,"normalizer_version":NORMALIZER_VERSION,
            "status":"queued","message":"A new immutable parser version was created; the source version remains unchanged."}


@app.delete("/api/parser-versions/{version_id}")
def delete_parser_version(version_id: int):
    with db() as conn:
        row=conn.execute("""SELECT v.*,d.project_id,d.active_parser_version_id FROM document_parser_versions v
          JOIN documents d ON d.id=v.document_id WHERE v.id=?""",(version_id,)).fetchone()
        if not row:raise HTTPException(404,"解析版本不存在")
        if row["is_active"] or row["active_parser_version_id"]==version_id:raise HTTPException(409,"active 解析版本不能删除；请先切换 active")
        if row["status"] in {"queued","submitted","running","visual_analyzing"}:raise HTTPException(409,"正在运行的解析版本不能删除")
        references=conn.execute("SELECT COUNT(*) FROM audit_run_documents WHERE parser_version_id=?",(version_id,)).fetchone()[0]
        if references:raise HTTPException(409,f"该解析版本已被 {references} 次审查运行固化引用，不能删除")
        children=conn.execute("SELECT COUNT(*) FROM document_parser_versions WHERE parent_version_id=?",(version_id,)).fetchone()[0]
        if children:raise HTTPException(409,"该解析版本仍有派生版本，不能删除")
        snapshot=dict(row)
    source=Path(row["artifact_dir"]).resolve() if row["artifact_dir"] else None;archived=None
    if source and source.exists():
        parser_root=PARSER_ARTIFACTS_DIR.resolve()
        if source!=parser_root and parser_root not in source.parents:raise HTTPException(409,"解析产物路径不在受管目录中，拒绝删除")
        trash=PARSER_ARTIFACTS_DIR.parent/"trash"/row["project_id"]
        trash.mkdir(parents=True,exist_ok=True)
        archived=(trash/f"{datetime.now().strftime('%Y%m%d-%H%M%S')}-parser-version-{version_id}").resolve()
        shutil.move(str(source),str(archived))
    try:
        with db() as conn:
            for table in ("visual_element_analyses","document_chunks","document_elements","document_parse_jobs"):
                conn.execute(f"DELETE FROM {table} WHERE parser_version_id=?",(version_id,))
            conn.execute("DELETE FROM document_parser_versions WHERE id=?",(version_id,))
            conn.execute("""INSERT INTO parser_version_deletions(version_id,document_id,project_id,snapshot_json,artifact_archive,deleted_at)
              VALUES(?,?,?,?,?,?)""",(version_id,row["document_id"],row["project_id"],json.dumps(snapshot,ensure_ascii=False),str(archived) if archived else None,datetime.now().strftime("%Y-%m-%d %H:%M:%S")))
    except Exception:
        if archived and archived.exists() and source and not source.exists():
            source.parent.mkdir(parents=True,exist_ok=True);shutil.move(str(archived),str(source))
        raise
    return {"ok":True,"version_id":version_id,"document_id":row["document_id"],"trace_run_id":row["trace_run_id"],
            "artifact_archived":str(archived) if archived else None,"recoverable":bool(archived),"trace_preserved":True}

@app.delete("/api/documents/{document_id}")
def delete_document(document_id: int):
    with db() as conn:
        row=conn.execute("SELECT id,project_id,name,file_path FROM documents WHERE id=?",(document_id,)).fetchone()
        if not row: raise HTTPException(404,"资料不存在")
        conn.execute("UPDATE documents SET deleted_at=?,status='已删除（历史证据保留）' WHERE id=?",(datetime.now().strftime("%Y-%m-%d %H:%M:%S"),document_id))
    source=Path(row["file_path"]) if row["file_path"] else None; archived=None
    if source and source.exists():
        trash=FILES_DIR.parent/"trash"/row["project_id"]
        trash.mkdir(parents=True,exist_ok=True)
        archived=trash/f"{datetime.now().strftime('%Y%m%d-%H%M%S')}-{source.name}"
        shutil.move(str(source),str(archived))
    artifact_source=PARSER_ARTIFACTS_DIR/str(document_id);artifact_archived=None
    if artifact_source.exists():
        trash=PARSER_ARTIFACTS_DIR.parent/"trash"/row["project_id"];trash.mkdir(parents=True,exist_ok=True)
        artifact_archived=trash/f"{datetime.now().strftime('%Y%m%d-%H%M%S')}-parser-{document_id}";shutil.move(str(artifact_source),str(artifact_archived))
    with db() as conn:
        if archived:conn.execute("UPDATE documents SET file_path=? WHERE id=?",(str(archived),document_id))
        if artifact_archived:
            for version in conn.execute("SELECT id,artifact_dir FROM document_parser_versions WHERE document_id=?",(document_id,)):
                if version["artifact_dir"]:
                    old=Path(version["artifact_dir"]);conn.execute("UPDATE document_parser_versions SET artifact_dir=? WHERE id=?",(str(artifact_archived/old.name),version["id"]))
    return {"ok":True,"document_id":document_id,"name":row["name"],"files_archived":str(archived) if archived else None,
            "parser_artifacts_archived":str(artifact_archived) if artifact_archived else None,"history_preserved":True}

@app.get("/api/ocr/batches/{batch_id}")
async def ocr_batch(batch_id: str):
    return await mineru_batch_result(batch_id)

@app.get("/api/documents/{document_id}/pages/{page_number}/render")
def render_document_page(document_id: int, page_number: int, scale: float=1.7):
    with db() as conn: row=conn.execute("SELECT file_path FROM documents WHERE id=?",(document_id,)).fetchone()
    if not row or not row["file_path"] or not Path(row["file_path"]).exists(): raise HTTPException(404,"原件不存在")
    path=Path(row["file_path"])
    if path.suffix.lower()!=".pdf": raise HTTPException(415,"当前仅支持 PDF 原页渲染")
    with fitz.open(path) as pdf:
        if page_number<1 or page_number>len(pdf): raise HTTPException(404,"页码不存在")
        pix=pdf[page_number-1].get_pixmap(matrix=fitz.Matrix(min(max(scale,.8),3),min(max(scale,.8),3)),alpha=False)
        return Response(pix.tobytes("png"),media_type="image/png",headers={"Cache-Control":"private, max-age=3600"})

@app.get("/api/documents/{document_id}/evidence-location")
def evidence_location(document_id: int, quote: str="", page: int=1, parser_version_id: int|None=None, element_id: str=""):
    with db() as conn:
        row=conn.execute("SELECT * FROM documents WHERE id=?",(document_id,)).fetchone()
        if not row: raise HTTPException(404,"资料不存在")
        item=decode(row,"content_json"); path=Path(item["file_path"]) if item.get("file_path") else None
        version=None
        if parser_version_id:
            version=conn.execute("SELECT * FROM document_parser_versions WHERE id=? AND document_id=?",(parser_version_id,document_id)).fetchone()
        elif item.get("active_parser_version_id"):
            version=conn.execute("SELECT * FROM document_parser_versions WHERE id=?",(item["active_parser_version_id"],)).fetchone()
        content=json.loads(version["content_json"] or "{}") if version else item["content_json"]
        if not version and path and path.exists() and path.suffix.lower()==".pdf" and not any(p.get("blocks") for p in content.get("pages") or []):
            pages,extractor=parse_document(path); content={**content,"pages":pages,"extractor":extractor,"layout_model":"paragraph-blocks"}
            conn.execute("UPDATE documents SET content_json=?,pages=?,status=? WHERE id=?",(json.dumps(content,ensure_ascii=False),len(pages),"已解析（版面块）",document_id))
        located=None
        if version and element_id:
            element=conn.execute("SELECT * FROM document_elements WHERE parser_version_id=? AND element_id=?",(version["id"],element_id)).fetchone()
            if element:
                meta=json.loads(element["metadata_json"] or "{}");located={"page":element["page"],"page_width":meta.get("page_width"),"page_height":meta.get("page_height"),
                    "block_id":element["element_id"],"bbox":json.loads(element["bbox_json"] or "null"),"block_text":element["markdown"] or element["text"],"element_type":element["element_type"],"html":element["html"],"cell_grid":json.loads(element["cell_grid_json"] or "[]")}
        if not located:located=locate_quote(content,quote,page)
        if located and path and path.exists() and path.suffix.lower()==".pdf" and (not located.get("page_width") or not located.get("page_height")):
            with fitz.open(path) as pdf:
                rect=pdf[int(located["page"])-1].rect; located["page_width"]=rect.width; located["page_height"]=rect.height
    page_number=int((located or {}).get("page") or page)
    return {"document_id":document_id,"parser_version_id":version["id"] if version else None,"page":page_number,"render_url":f"/api/documents/{document_id}/pages/{page_number}/render","location":located,"coordinate_source":content.get("extractor"),"is_original_page":bool(path and path.suffix.lower()==".pdf")}

class ProviderTest(BaseModel):
    provider_id: str | None = None
    provider: dict | None = None

@app.get("/api/providers")
def providers():
    config=public_config()
    return {"items":config["providers"],"providers":config["providers"],"stage_routes":config["stage_routes"],
            "schema_version":config["schema_version"],"storage":"backend/data/model-config.json",
            "warning":"MVP 明文保存供应商密钥；API 与 Trace 已脱敏"}

@app.get("/api/providers/route-templates")
def provider_route_templates():
    """Distinct model-route snapshots from completed immutable audit runs."""
    current_ids={p["id"] for p in load_config()["providers"]};templates=[];seen=set()
    with db() as conn:
        history=conn.execute("""SELECT r.id,r.project_id,r.project_run_no,r.started_at,p.name AS project_name,r.config_snapshot_json
          FROM audit_runs r JOIN projects p ON p.id=r.project_id
          WHERE r.run_kind='audit' AND r.status='completed' ORDER BY r.id DESC""").fetchall()
    for row in history:
        try:snapshot_data=json.loads(row["config_snapshot_json"] or "{}")
        except Exception:continue
        routes=snapshot_data.get("stage_routes") or {}
        if not routes:continue
        signature=json.dumps(routes,ensure_ascii=False,sort_keys=True,separators=(",",":"))
        if signature in seen:continue
        seen.add(signature)
        names={p.get("id"):p.get("name") for p in snapshot_data.get("providers",[]) if isinstance(p,dict)}
        missing=sorted({str(v.get("provider_id")) for v in routes.values() if isinstance(v,dict) and v.get("provider_id") not in current_ids})
        templates.append({"id":f"run-{row['id']}","source_run_id":row["id"],"project_id":row["project_id"],
          "project_name":row["project_name"],"project_run_no":row["project_run_no"],"started_at":row["started_at"],
          "stage_routes":routes,"provider_names":names,"missing_provider_ids":missing,"applicable":not missing})
        if len(templates)>=20:break
    return {"items":templates,"total":len(templates),"source":"completed immutable audit snapshots"}

@app.post("/api/providers/test")
async def providers_test(body: ProviderTest):
    if body.provider:
        candidate=dict(body.provider)
        if not candidate.get("api_key") and body.provider_id:
            existing=load_config();old=next((p for p in existing["providers"] if p["id"]==body.provider_id),None)
            if old:candidate["api_key"]=old.get("api_key","");candidate["api_key_env"]=old.get("api_key_env","")
        result=await test_provider_config(candidate)
    elif body.provider_id:result=await test_provider(body.provider_id)
    else:raise HTTPException(400,"请提供 provider_id 或待测试的 provider 配置")
    if result.get("status_code")==404: raise HTTPException(404,result["message"])
    return result


class ModelConfigBody(BaseModel):
    schema_version: int = 1
    providers: list[dict]
    stage_routes: dict


@app.put("/api/providers")
def providers_update(body: ModelConfigBody):
    try:
        saved=merge_public_update(body.model_dump())
    except ValueError as exc:
        raise HTTPException(400,str(exc))
    return {"ok":True,"config":public_config(saved),"warning":"密钥以明文保存在本机配置文件；仅限可信内网"}


@app.get("/api/providers/{provider_id}/models")
async def provider_models(provider_id: str):
    try: return await discover_models(provider_id)
    except ValueError as exc: raise HTTPException(400,str(exc))
    except Exception as exc: raise HTTPException(502,str(exc)[:300])


@app.post("/api/demo/install")
def install_demo_projects():
    with db() as conn: before=conn.execute("SELECT COUNT(*) FROM projects WHERE is_demo=1").fetchone()[0]
    seed()
    with db() as conn: after=conn.execute("SELECT COUNT(*) FROM projects WHERE is_demo=1").fetchone()[0]
    return {"ok":True,"installed":after-before,"demo_projects":after,"message":"示例项目已隔离，不计入统计或风险角标"}


@app.get("/api/prompts")
def prompts(): return {"items":list_prompts()}


class PromptDraftBody(BaseModel):
    stage: str
    based_on_id: int | None = None


@app.post("/api/prompts/drafts",status_code=201)
def prompt_draft(body: PromptDraftBody):
    try: return create_draft(body.stage,body.based_on_id)
    except KeyError: raise HTTPException(404,"Prompt 不存在")


@app.patch("/api/prompts/{prompt_id}")
def prompt_update(prompt_id: int, body: dict):
    try: return update_draft(prompt_id,body)
    except KeyError: raise HTTPException(404,"Prompt 不存在")
    except ValueError as exc: raise HTTPException(409,str(exc))


@app.post("/api/prompts/{prompt_id}/publish")
def prompt_publish(prompt_id: int):
    try: return publish(prompt_id)
    except KeyError: raise HTTPException(404,"Prompt 不存在")
    except ValueError as exc: raise HTTPException(409,str(exc))


@app.get("/api/trace/projects")
def trace_projects():
    with db() as conn:
        items=rows(conn.execute("""SELECT p.id,p.name,p.community,p.is_demo,COUNT(r.id) AS run_count,MAX(r.id) AS latest_run_id,
          MAX(r.started_at) AS latest_started_at FROM projects p LEFT JOIN audit_runs r ON r.project_id=p.id
          GROUP BY p.id ORDER BY latest_run_id DESC,p.updated_at DESC""").fetchall())
    return {"items":items}


@app.get("/api/trace/projects/{project_id}/runs")
def trace_project_runs(project_id: str):
    with db() as conn:
        if not conn.execute("SELECT 1 FROM projects WHERE id=?",(project_id,)).fetchone(): raise HTTPException(404,"项目不存在")
        items=rows(conn.execute("""SELECT r.id,r.project_id,r.parent_run_id,r.derived_from_stage,r.run_kind,r.started_at,r.finished_at,r.status,
          r.fact_count,r.anomaly_count,r.provider,r.progress,r.current_stage,r.error,r.project_run_no
          FROM audit_runs r WHERE r.project_id=? ORDER BY r.id DESC""",(project_id,)).fetchall())
    return {"items":items}


@app.get("/api/trace/runs/{run_id}")
def trace_detail(run_id: int):
    try: return run_trace(run_id,True)
    except KeyError: raise HTTPException(404,"运行不存在")


@app.get("/api/trace/runs/{run_id}/export")
def trace_export(run_id: int):
    try: payload=run_trace(run_id,True)
    except KeyError: raise HTTPException(404,"运行不存在")
    with db() as conn:
        payload["documents"]=[{"id":r["id"],"name":r["name"],"doc_type":r["doc_type"],"sha256":r["sha256"],
          "pages":r["pages"],"download_url":f"/api/documents/{r['id']}/download"} for r in conn.execute(
          "SELECT id,name,doc_type,sha256,pages FROM documents WHERE project_id=? AND deleted_at IS NULL ORDER BY id",(payload["run"]["project_id"],))]
    return JSONResponse(payload,headers={"Content-Disposition":f'attachment; filename="trace-run-{run_id}.json"'})


@app.get("/api/trace/runs/{run_id}/bundle")
def trace_bundle(run_id: int):
    try: payload=run_trace(run_id,True)
    except KeyError: raise HTTPException(404,"运行不存在")
    with db() as conn:
        docs=[dict(r) for r in conn.execute("SELECT id,name,sha256,file_path FROM documents WHERE project_id=? AND deleted_at IS NULL ORDER BY id",(payload["run"]["project_id"],))]
    payload["documents"]=[{"id":d["id"],"name":d["name"],"sha256":d["sha256"],"archive_path":f"originals/{d['id']}-{Path(d['name']).name}"} for d in docs]
    export_dir=TRACES_DIR/"exports";export_dir.mkdir(parents=True,exist_ok=True)
    handle=tempfile.NamedTemporaryFile(prefix=f"trace-run-{run_id}-",suffix=".zip",dir=export_dir,delete=False); path=Path(handle.name);handle.close()
    with zipfile.ZipFile(path,"w",zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("trace.json",json.dumps(payload,ensure_ascii=False,indent=2))
        for document in docs:
            source=Path(document["file_path"] or "")
            if source.is_file(): archive.write(source,f"originals/{document['id']}-{Path(document['name']).name}")
    return FileResponse(path,media_type="application/zip",filename=f"trace-run-{run_id}-with-originals.zip",
                        background=BackgroundTask(lambda: path.unlink(missing_ok=True)))


def _recalculate_project_after_run_delete(conn, project_id: str):
    latest=conn.execute("SELECT id FROM audit_runs WHERE project_id=? AND status='completed' AND run_kind='audit' ORDER BY id DESC LIMIT 1",(project_id,)).fetchone()
    if not latest:
        conn.execute("UPDATE projects SET risk_count=0,high_risk_count=0,facts_json='{}',updated_at=? WHERE id=?",
                     (datetime.now().strftime("%Y-%m-%d %H:%M"),project_id))
        return None
    pending=conn.execute("SELECT COUNT(*) FROM anomalies WHERE run_id=? AND status='待复核'",(latest["id"],)).fetchone()[0]
    high=conn.execute("SELECT COUNT(*) FROM anomalies WHERE run_id=? AND status='待复核' AND severity='high'",(latest["id"],)).fetchone()[0]
    facts={}
    for row in conn.execute("SELECT field_name,value_json FROM extracted_facts WHERE run_id=? ORDER BY confidence DESC,id",(latest["id"],)):
        if row["field_name"] not in facts:
            try: facts[row["field_name"]]=json.loads(row["value_json"])
            except Exception: pass
    conn.execute("UPDATE projects SET risk_count=?,high_risk_count=?,facts_json=?,updated_at=? WHERE id=?",
                 (pending,high,json.dumps(facts,ensure_ascii=False),datetime.now().strftime("%Y-%m-%d %H:%M"),project_id))
    return latest["id"]


@app.delete("/api/trace/runs/{run_id}")
def trace_delete(run_id: int, confirm: str=""):
    if confirm!=f"run-{run_id}": raise HTTPException(400,f"删除确认值必须为 run-{run_id}")
    with db() as conn:
        run=conn.execute("SELECT project_id,action_id FROM audit_runs WHERE id=?",(run_id,)).fetchone()
        if not run: raise HTTPException(404,"运行不存在")
        paths=[Path(r[0]) for r in conn.execute("SELECT file_path FROM trace_artifacts WHERE run_id=? AND file_path IS NOT NULL",(run_id,))]
        conn.execute("UPDATE audit_runs SET parent_run_id=NULL WHERE parent_run_id=?",(run_id,))
        conn.execute("UPDATE audit_runs SET baseline_run_id=NULL WHERE baseline_run_id=?",(run_id,))
        conn.execute("UPDATE document_parser_versions SET trace_run_id=NULL WHERE trace_run_id=?",(run_id,))
        conn.execute("DELETE FROM audit_run_documents WHERE run_id=?",(run_id,))
        anomaly_ids=[row[0] for row in conn.execute("SELECT id FROM anomalies WHERE run_id=?",(run_id,))]
        if anomaly_ids:
            placeholders=','.join('?' for _ in anomaly_ids)
            conn.execute(f"UPDATE anomalies SET prior_anomaly_id=NULL WHERE prior_anomaly_id IN ({placeholders})",anomaly_ids)
            conn.execute(f"UPDATE risk_transitions SET prior_anomaly_id=NULL WHERE prior_anomaly_id IN ({placeholders})",anomaly_ids)
            conn.execute(f"UPDATE risk_transitions SET current_anomaly_id=NULL WHERE current_anomaly_id IN ({placeholders})",anomaly_ids)
        conn.execute("DELETE FROM risk_transitions WHERE run_id=?",(run_id,))
        for table in ("audit_events","ai_calls","extracted_facts","anomalies"):
            conn.execute(f"DELETE FROM {table} WHERE run_id=?",(run_id,))
        conn.execute("UPDATE run_spans SET input_artifact_id=NULL,output_artifact_id=NULL,error_artifact_id=NULL WHERE run_id=?",(run_id,))
        conn.execute("DELETE FROM trace_artifacts WHERE run_id=?",(run_id,))
        conn.execute("DELETE FROM run_spans WHERE run_id=?",(run_id,))
        conn.execute("DELETE FROM audit_runs WHERE id=?",(run_id,))
        if run["action_id"]: conn.execute("DELETE FROM review_actions WHERE id=?",(run["action_id"],))
        reverted_to=_recalculate_project_after_run_delete(conn,run["project_id"])
    for path in paths:
        if path.exists(): path.unlink()
    run_dir=TRACES_DIR/f"run-{run_id}"
    if run_dir.exists() and not any(run_dir.iterdir()): run_dir.rmdir()
    return {"ok":True,"deleted_run_id":run_id,"project_id":run["project_id"],"reverted_to_run_id":reverted_to}


class ReplayBody(BaseModel):
    stage: str
    mode: str = "probe"
    route_overrides: dict | None = None
    prompt_id: int | None = None


@app.post("/api/trace/runs/{run_id}/replay",status_code=201)
async def trace_replay(run_id: int, body: ReplayBody, background_tasks: BackgroundTasks):
    if body.stage not in ALL_STAGES: raise HTTPException(400,"未知阶段")
    with db() as conn:
        source=conn.execute("SELECT * FROM audit_runs WHERE id=?",(run_id,)).fetchone()
        if not source: raise HTTPException(404,"源运行不存在")
    if body.mode=="derived":
        source_trace=run_trace(run_id,True)
        source_artifacts={a["id"]:a for a in source_trace["artifacts"]}
        stage_order=["ocr","document_classification","general_extraction","payment_extraction","change_extraction",
                     "contract_extraction","semantic_consistency","policy_retrieval","text2sql","result_summary"]
        if body.stage not in stage_order:
            raise HTTPException(409,"该阶段支持单阶段测试，但不在当前派生模型链中；请选择受支持的审查阶段")
        start_index=stage_order.index(body.stage)
        prompts_now=prompt_snapshot(); config_now=snapshot(body.route_overrides or {})
        started=datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        with db() as conn:
            cur=conn.execute("""INSERT INTO audit_runs(project_id,started_at,status,rule_count,fact_count,anomaly_count,provider,progress,current_stage,result_json,
              parent_run_id,derived_from_stage,run_kind,config_snapshot_json,prompt_versions_json,route_overrides_json)
              VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",(source["project_id"],started,"running",0,0,0,"derived-trace",0,body.stage,"{}",run_id,body.stage,
              "derived",json.dumps(config_now,ensure_ascii=False),json.dumps(prompts_now,ensure_ascii=False),json.dumps(body.route_overrides or {},ensure_ascii=False)))
            derived_run=int(cur.lastrowid)
        outputs={}
        try:
            for index,stage in enumerate(stage_order):
                source_calls=[s for s in source_trace["spans"] if s["stage"]==stage and s["kind"]=="model_call"]
                if index<start_index:
                    reused=start_span(derived_run,stage,"reused_snapshot",f"复用 Run #{run_id} 的上游快照")
                    attach_input(reused,{"source_run_id":run_id,"source_spans":[s["id"] for s in source_calls],
                        "inputs":[source_artifacts.get(s.get("input_artifact_id"),{}).get("content") for s in source_calls]})
                    finish_span(reused,"reused",output={"source_run_id":run_id,"outputs":[source_artifacts.get(s.get("output_artifact_id"),{}).get("content") for s in source_calls]},
                                metadata={"immutable_source":True,"reused_upstream":True})
                    continue
                if stage=="ocr":
                    skipped_span(derived_run,stage,"OCR 派生运行不重复提交外部解析任务；复用源运行文档快照")
                    continue
                if not source_calls:
                    skipped_span(derived_run,stage,"源运行该阶段没有模型调用")
                    continue
                prompt=published_prompt(stage)
                stage_outputs=[]
                for call in source_calls:
                    artifact=source_artifacts.get(call.get("input_artifact_id"))
                    if not artifact: continue
                    request=artifact["content"];variables=request.get("prompt_variables") or {}
                    messages=[{"role":"system","content":prompt["system_prompt"]},
                              {"role":"user","content":render(prompt["user_prompt"],variables)}]
                    parsed=await model_json(derived_run,stage,messages,max_tokens=int(request.get("max_tokens") or 2400),
                        route_overrides=body.route_overrides or {},prompt_version=prompt,prompt_variables=variables)
                    parsed.pop("__trace",None);stage_outputs.append(parsed)
                outputs[stage]=stage_outputs
                with db() as conn: conn.execute("UPDATE audit_runs SET current_stage=?,progress=? WHERE id=?",
                    (stage,round((index+1)/len(stage_order)*100),derived_run))
            with db() as conn: conn.execute("UPDATE audit_runs SET status='completed',finished_at=?,progress=100,current_stage='complete',result_json=? WHERE id=?",
                (datetime.now().strftime("%Y-%m-%d %H:%M:%S"),json.dumps({"source_run_id":run_id,"replay_from":body.stage,
                 "upstream":"reused immutable source snapshots","downstream_outputs":outputs,"project_mutated":False},ensure_ascii=False),derived_run))
        except Exception as exc:
            with db() as conn: conn.execute("UPDATE audit_runs SET status='failed',finished_at=?,error=?,result_json=? WHERE id=?",
                (datetime.now().strftime("%Y-%m-%d %H:%M:%S"),str(exc),json.dumps({"source_run_id":run_id,"partial_outputs":outputs},ensure_ascii=False),derived_run))
        return {"ok":True,"run_id":derived_run,"mode":"derived","source_run_id":run_id,"stage":body.stage,
                "project_mutated":False}
    if body.mode!="probe": raise HTTPException(400,"mode 仅支持 probe 或 derived")
    with db() as conn:
        span=conn.execute("SELECT input_artifact_id FROM run_spans WHERE run_id=? AND stage=? AND kind='model_call' AND input_artifact_id IS NOT NULL ORDER BY sequence LIMIT 1",(run_id,body.stage)).fetchone()
        if not span: raise HTTPException(409,"该阶段没有可重放的模型输入")
        source_prompts=json.loads(source["prompt_versions_json"] or "{}")
        started=datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        cur=conn.execute("""INSERT INTO audit_runs(project_id,started_at,status,rule_count,fact_count,anomaly_count,provider,progress,current_stage,result_json,
          parent_run_id,derived_from_stage,run_kind,config_snapshot_json,prompt_versions_json,route_overrides_json)
          VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",(source["project_id"],started,"running",0,0,0,"prompt-test",0,body.stage,"{}",run_id,body.stage,
          "prompt_test",json.dumps(snapshot(body.route_overrides or {}),ensure_ascii=False),source["prompt_versions_json"],json.dumps(body.route_overrides or {},ensure_ascii=False)))
        test_run=int(cur.lastrowid)
    artifact=load_artifact(int(span["input_artifact_id"]))
    request=artifact["content"]
    selected_prompt=next((p for p in list_prompts() if p["id"]==body.prompt_id),None) if body.prompt_id else source_prompts.get(body.stage)
    if body.prompt_id and (not selected_prompt or selected_prompt["stage"]!=body.stage): raise HTTPException(400,"Prompt 与重放阶段不匹配")
    messages=request["messages"]
    variables=request.get("prompt_variables") or {}
    if selected_prompt:
        messages=[{"role":"system","content":selected_prompt["system_prompt"]},
                  {"role":"user","content":render(selected_prompt["user_prompt"],variables)}]
    try:
        output=await model_json(test_run,body.stage,messages,max_tokens=int(request.get("max_tokens") or 2400),
            route_overrides=body.route_overrides or {},prompt_version=selected_prompt,prompt_variables=variables)
        output.pop("__trace",None)
        with db() as conn: conn.execute("UPDATE audit_runs SET status='completed',finished_at=?,progress=100,result_json=? WHERE id=?",
            (datetime.now().strftime("%Y-%m-%d %H:%M:%S"),json.dumps({"probe_output":output},ensure_ascii=False),test_run))
    except Exception as exc:
        with db() as conn: conn.execute("UPDATE audit_runs SET status='failed',finished_at=?,error=? WHERE id=?",
            (datetime.now().strftime("%Y-%m-%d %H:%M:%S"),str(exc),test_run))
    return {"ok":True,"run_id":test_run,"mode":"probe","source_run_id":run_id,"stage":body.stage}
